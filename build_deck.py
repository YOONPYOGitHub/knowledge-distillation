# -*- coding: utf-8 -*-
"""
eval_section_deck.md  ->  eval_section_deck.pptx
ppt_guideline_grad.md 규칙 반영: 폰트 하한 / Navy 팔레트 / noAutofit+wrap /
§6-4 컨텐츠·여백 밸런스(동적 카드 높이 + 세로 중앙) / §6-5 텍스트 넘침 사전 점검.
"""
import sys, re, os, math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
# 사용법: python build_deck.py <입력.md> [출력.pptx]
#   - 인자가 없으면 같은 폴더의 기본 예시(eval_section_deck.md)를 변환
if len(sys.argv) > 1:
    SRC = os.path.abspath(sys.argv[1])
    OUT = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else os.path.splitext(SRC)[0] + ".pptx"
else:
    print("사용법: python build_deck.py <입력.md> [출력.pptx]")
    SRC = os.path.join(BASE, "eval_section_deck.md")
    OUT = os.path.splitext(SRC)[0] + ".pptx"
    print(f"  (인자 없음 → 기본 예시 변환: {os.path.basename(SRC)})\n")

FONT   = "Pretendard"
NAVY   = RGBColor(0x1B, 0x2A, 0x4E); BLUE  = RGBColor(0x2E, 0x6F, 0xD4)
ACCENT = RGBColor(0xC0, 0x39, 0x2B); TEXT  = RGBColor(0x22, 0x22, 0x22)
CARD   = RGBColor(0xF8, 0xFA, 0xFD); STRIP = RGBColor(0xEE, 0xF3, 0xFB)
LINE   = RGBColor(0xD6, 0xDE, 0xEA); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FLOOR  = {"headline": 28, "body": 17, "sub": 15, "footnote": 11}

HEAD_PT, TITLE_PT, BODY_PT = 30, 22, 20      # 시니어 발표 + 카드 채움
PAGE_W, PAGE_H = 13.333, 7.5
MX, BODY_TOP, PAGE_BOT, E_H = 0.5, 1.80, 6.95, 0.82

used, report = [], []

# ---------- low-level ----------
def _set_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)

def _no_autofit(tf):
    bp = tf._txBody.find(qn("a:bodyPr"))
    for t in ("a:noAutofit", "a:normAutofit", "a:spAutoFit"):
        e = bp.find(qn(t))
        if e is not None: bp.remove(e)
    bp.append(bp.makeelement(qn("a:noAutofit"), {}))
    bp.set("wrap", "square"); tf.word_wrap = True

def _style(run, size, color=TEXT, bold=False, role="body"):
    run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    _set_font(run); used.append((role, size))

def add_text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.18):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; _no_autofit(tf); tf.vertical_anchor = anchor
    tf.margin_left = Emu(45000); tf.margin_right = Emu(45000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(5)
        for txt, opt in para:
            r = p.add_run(); r.text = txt
            _style(r, opt.get("size", BODY_PT), opt.get("color", TEXT),
                   opt.get("bold", False), opt.get("role", "body"))
    return tb

def add_rect(slide, x, y, w, h, fill, line=None, rounded=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s

# ---------- §6-5 텍스트 줄 수 / 박스 높이 예측 ----------
def _text_w(text, pt):
    return sum(pt * (1.0 if ord(c) >= 0xAC00 else 0.6) for c in text)

def est_lines(text, pt, box_w_in, pad=0.20):
    avail = (box_w_in - 2 * pad) * 72
    return max(1, math.ceil(_text_w(text, pt) / avail)) if avail > 0 else 1

def group_height(g, w_in, title_pt=TITLE_PT):
    """제목 줄수 + 불릿 줄수로 카드 필요 높이(in) 산정 (§6-4 동적 높이)."""
    h = 0.18 * 2
    h += est_lines(g["title"], title_pt, w_in) * (title_pt * 1.25 / 72) + 0.12
    for b in g["bullets"]:
        h += est_lines("•  " + b, BODY_PT, w_in) * (BODY_PT * 1.40 / 72) + 0.08
    return h

def clamp(v, lo, hi):
    return max(lo, min(v, hi))

# ---------- components ----------
def render_group(slide, x, y, w, h, g, title_pt=TITLE_PT):
    add_rect(slide, x, y, w, h, CARD, line=LINE, rounded=True)
    paras = [[(g["title"], {"size": title_pt, "color": NAVY, "bold": True})]]
    for b in g["bullets"]:
        paras.append([("•  ", {"size": BODY_PT, "color": BLUE, "bold": True}),
                      (b, {"size": BODY_PT, "color": TEXT})])
    add_text(slide, x + 0.22, y + 0.10, w - 0.44, h - 0.20, paras,
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.22)   # 세로 중앙

def emphasis_bar(slide, x, y, w, h, text):
    add_rect(slide, x, y, w, h, STRIP, rounded=True)
    add_rect(slide, x, y, 0.10, h, NAVY)
    color = ACCENT if text.strip().startswith("→") else NAVY
    add_text(slide, x + 0.30, y, w - 0.5, h,
             [[(text, {"size": 18, "color": color, "bold": True})]],
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)

def render_table(slide, x, y, w, rows):
    nr, nc = len(rows), len(rows[0])
    tbl = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(0.46 * nr)).table
    if nc == 4:
        for i, cw in enumerate([3.0, 1.3, 4.4, 3.63]):
            tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            c.margin_left = Emu(70000); c.margin_right = Emu(50000)
            c.margin_top = Emu(28000); c.margin_bottom = Emu(28000)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE; c.fill.solid()
            if ri == 0:
                c.fill.fore_color.rgb = NAVY; col, bold, sz = WHITE, True, 16
            else:
                c.fill.fore_color.rgb = WHITE if ri % 2 else CARD
                col, bold, sz = TEXT, (ci == 0), 15
            _no_autofit(c.text_frame)
            p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(val)
            _style(r, sz, col, bold, "sub")
    return 0.46 * nr

# ---------- parse ----------
def parse_slides(md):
    pat = re.compile(r"<!-- slide:(\d+) \| layout:(\S+) \| motif:(\S+) -->(.*?)(?=\n---\n<!-- slide:|\Z)", re.S)
    out = []
    for num, layout, motif, body in pat.findall(md):
        head = None; groups = []; cur = None; table = []; emph = []; notes = []
        for raw in body.splitlines():
            s = raw.strip()
            if not s: continue
            if s.startswith("## "):
                head = s[3:].strip()
            elif s.startswith("|") and s.endswith("|"):
                cells = [c.strip() for c in s.strip("|").split("|")]
                if set("".join(cells)) <= set("-: "): continue
                table.append(cells)
            elif re.fullmatch(r"\*\*(.+?)\*\*", s):
                inner = s[2:-2].strip()
                if inner.startswith("→") or len(inner) > 42:
                    emph.append(inner)
                else:
                    cur = {"title": inner, "bullets": []}; groups.append(cur)
            elif s.startswith("**") and "**:" in s:
                emph.append(s.replace("**", ""))
            elif s.startswith("- "):
                if cur is not None: cur["bullets"].append(s[2:].strip())
            elif s.startswith("> 발표노트:") or s.startswith("> 출처:"):
                notes.append(s[2:].strip())
        out.append(dict(num=int(num), layout=layout, motif=motif, head=head,
                        groups=groups, table=table, emph=emph, notes=notes))
    return out

# ---------- build ----------
def build():
    md = open(SRC, encoding="utf-8").read()
    slides = parse_slides(md)
    prs = Presentation(); prs.slide_width = Inches(PAGE_W); prs.slide_height = Inches(PAGE_H)
    blank = prs.slide_layouts[6]

    for sd in slides:
        slide = prs.slides.add_slide(blank)
        add_text(slide, MX, 0.34, PAGE_W - 2 * MX, 1.20,
                 [[(sd["head"], {"size": HEAD_PT, "color": NAVY, "bold": True, "role": "headline"})]],
                 spacing=1.06)
        add_rect(slide, MX, 1.66, PAGE_W - 2 * MX, 0.028, NAVY)
        emph = sd["emph"]; ne = len(emph)

        if sd["layout"] in ("three-pillar", "two-column"):
            n = len(sd["groups"]); gap = 0.25 if n >= 3 else 0.30
            w = (PAGE_W - 2 * MX - gap * (n - 1)) / n
            tpt = 19 if sd["layout"] == "three-pillar" else TITLE_PT
            bottom = PAGE_BOT - (ne * E_H + 0.15) if ne else PAGE_BOT
            area = bottom - BODY_TOP
            need = max(group_height(g, w, tpt) for g in sd["groups"])
            H = clamp(need * 1.06, area * 0.46, area * 0.92)   # 동적 + 영역 비례 clamp
            cy = BODY_TOP + (area - H) / 2.0                   # 세로 중앙
            for i, g in enumerate(sd["groups"]):
                render_group(slide, MX + i * (w + gap), cy, w, H, g, title_pt=tpt)
            ey = PAGE_BOT - ne * E_H
            for e in emph:
                emphasis_bar(slide, MX, ey, PAGE_W - 2 * MX, E_H - 0.06, e); ey += E_H
            report.append((sd["num"], sd["layout"], H, area, min(need, H) / H * 100))

        elif sd["layout"] == "appendix-table":
            th = render_table(slide, MX, BODY_TOP + 0.05, PAGE_W - 2 * MX, sd["table"])
            ey = BODY_TOP + 0.05 + th + 0.28
            for e in emph:
                emphasis_bar(slide, MX, ey, PAGE_W - 2 * MX, E_H - 0.06, e); ey += E_H
            area = PAGE_BOT - BODY_TOP
            report.append((sd["num"], sd["layout"], th, area, th / area * 100))

        else:
            paras = []
            for g in sd["groups"]:
                paras.append([(g["title"], {"size": TITLE_PT, "color": NAVY, "bold": True})])
                for b in g["bullets"]:
                    paras.append([("•  ", {"size": BODY_PT, "color": BLUE, "bold": True}),
                                  (b, {"size": BODY_PT, "color": TEXT})])
            add_text(slide, MX + 0.1, BODY_TOP + 0.1, PAGE_W - 2 * MX - 0.2, 4.6, paras)
            report.append((sd["num"], sd["layout"], 0, 0, 0))

        if sd["notes"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(sd["notes"])

    prs.save(OUT)
    return slides

# ---------- 자체 검증 ----------
def verify(slides):
    roles = {}
    for r, p in used:
        roles.setdefault(r, []).append(p)
    hmin = min(roles.get("headline", [99])); bmin = min(roles.get("body", [99])); smin = min(roles.get("sub", [99]))
    ok = hmin >= FLOOR["headline"] and bmin >= FLOOR["body"] and smin >= FLOOR["sub"]
    print(f"슬라이드 {len(slides)}장 -> {OUT}")
    print(f"폰트: headline {hmin}pt / body {bmin}pt / sub·table {smin}pt | 하한 28/17/15 -> {'PASS' if ok else 'FAIL'}")
    print("레이아웃 밸런스 (카드높이 H / 본문영역 area / 카드 내 채움%):")
    for num, lay, H, area, fill in report:
        flag = "" if lay == "appendix-table" else ("  ← 채움 낮음" if fill < 65 else "")
        print(f"  slide {num:>2} {lay:<14} H={H:>4.2f}in area={area:>4.2f}in fill={fill:>4.0f}%{flag}")

if __name__ == "__main__":
    verify(build())
